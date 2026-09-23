"""Gerador HAScript MOVFIN integrado ao Acerto RSC.
ODS usa o leitor local. XLSX/XLSM exigem openpyxl.
Mantenha este arquivo junto de extrair_rsc.py.
"""

from __future__ import annotations

import html
import re
import threading
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Dict, List, Tuple

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

COLUNAS_OBRIGATORIAS = [
    "SIAPE/7", "Nome", "MÊS/24", "Rubrica", "R/D", "SEQ",
    "Valor/0000,00", "Justificativa", "DOC. Legal",
]

COLUNAS_EXIGIDAS_PARA_GERAR = [
    "SIAPE/7", "MÊS/24", "Rubrica", "R/D", "SEQ",
    "Valor/0000,00", "Justificativa", "DOC. Legal",
]


def texto(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def normalizar_nome_coluna(v: Any) -> str:
    nome = texto(v).replace("\n", " ").replace("  ", " ").strip()
    return {"Valor/0000": "Valor/0000,00", "Doc legal": "DOC. Legal"}.get(nome, nome)


def xml_attr(v: Any) -> str:
    return html.escape(texto(v), quote=True)


def somente_digitos(v: Any) -> str:
    return re.sub(r"\D", "", texto(v).split(".")[0])


def formatar_siape(v: Any) -> str:
    return somente_digitos(v)


def siape_valido(v: Any) -> bool:
    s = formatar_siape(v)
    return s.isdigit() and 6 <= len(s) <= 8


def formatar_rubrica(v: Any) -> str:
    digitos = somente_digitos(v)
    return digitos.zfill(5) if digitos else texto(v)


def formatar_seq(v: Any) -> str:
    return somente_digitos(v) or texto(v).split(".")[0]


def formatar_valor(v: Any) -> str:
    if v is None or texto(v) == "":
        return ""
    if isinstance(v, (int, float, Decimal)):
        d = Decimal(str(v))
    else:
        s = texto(v).replace("R$", "").strip()
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        try:
            d = Decimal(s)
        except InvalidOperation:
            return texto(v)
    d = d.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"{d:.2f}".replace(".", ",")


def formatar_mes(v: Any) -> str:
    s = texto(v).replace(" ", "")
    if not s:
        return ""
    m = re.match(r"^([A-Za-zÀ-ÿ]{3})/(\d{2}|\d{4})$", s)
    if m:
        mes, ano = m.groups()
        return f"{mes}{'20' + ano if len(ano) == 2 else ano}"
    m = re.match(r"^([A-Za-zÀ-ÿ]{3})(\d{2}|\d{4})$", s)
    if m:
        mes, ano = m.groups()
        return f"{mes}{'20' + ano if len(ano) == 2 else ano}"
    return s


def localizar_cabecalho_em_matriz(matriz: List[List[Any]]) -> Tuple[int, Dict[str, int]]:
    for idx, row in enumerate(matriz):
        valores = [normalizar_nome_coluna(c) for c in row]
        if "SIAPE/7" in valores and "Rubrica" in valores:
            mapa = {nome: i for i, nome in enumerate(valores) if nome}
            faltando = [c for c in COLUNAS_OBRIGATORIAS if c not in mapa]
            if faltando:
                raise ValueError("Cabeçalho encontrado, mas faltam colunas: " + ", ".join(faltando))
            return idx, mapa
    raise ValueError('Não encontrei o cabeçalho. A aba precisa ter as colunas "SIAPE/7" e "Rubrica".')


def ler_entrada_matriz(matriz: List[List[Any]]) -> Tuple[List[Dict[str, Any]], List[str]]:
    linha_header, mapa = localizar_cabecalho_em_matriz(matriz)
    registros: List[Dict[str, Any]] = []
    ignoradas: List[str] = []

    for idx in range(linha_header + 1, len(matriz)):
        row = matriz[idx]
        item: Dict[str, Any] = {}
        for col in COLUNAS_OBRIGATORIAS:
            pos = mapa[col]
            item[col] = row[pos] if pos < len(row) else None

        if not any(texto(v) for v in item.values()):
            continue

        linha_real = idx + 1
        if not siape_valido(item["SIAPE/7"]):
            ignoradas.append(f"Linha {linha_real}: ignorada, SIAPE inválido ou vazio")
            continue

        faltando = [col for col in COLUNAS_EXIGIDAS_PARA_GERAR if not texto(item[col])]
        if faltando:
            ignoradas.append(f"Linha {linha_real}: ignorada, faltando {', '.join(faltando)}")
            continue

        item["__linha_excel__"] = linha_real
        registros.append(item)

    if not registros:
        raise ValueError("Nenhuma linha válida foi encontrada abaixo do cabeçalho.")
    return registros, ignoradas


def ler_xlsx_xlsm(caminho: Path, aba: str) -> Tuple[List[Dict[str, Any]], List[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(caminho, data_only=True, keep_vba=caminho.suffix.lower() == ".xlsm")
    if aba not in wb.sheetnames:
        raise ValueError(f'Aba "{aba}" não encontrada. Abas disponíveis: {", ".join(wb.sheetnames)}')
    ws = wb[aba]
    matriz = [[cell.value for cell in row] for row in ws.iter_rows()]
    return ler_entrada_matriz(matriz)


def ler_ods(caminho: Path, aba: str):
    # Reutiliza o leitor ODS do Acerto RSC, sem exigir odfpy.
    import extrair_rsc as leitor
    _, abas = leitor.abrir(caminho)
    tabela = leitor.escolher(abas, aba)
    matriz = []
    for linha in leitor.linhas(tabela):
        valores = []
        for cell in leitor.elementos(linha, {'table-cell', 'covered-table-cell'}):
            valor = leitor.valor(cell)
            if isinstance(valor, tuple) and valor[0] == 'date':
                valor = valor[1]
            repeticoes = leitor.repeticoes(cell, 'number-columns-repeated')
            valores.extend([valor] * min(repeticoes, max(0, 100-len(valores))))
        repeticoes = leitor.repeticoes(linha, 'number-rows-repeated')
        if not any(texto(v) for v in valores):
            # Linhas vazias finais compactadas não afetam a leitura dos registros.
            if repeticoes > 1000:
                continue
        if len(matriz) + repeticoes > 1048576:
            raise ValueError('Quantidade de linhas excede o limite suportado.')
        for _ in range(repeticoes):
            matriz.append(list(valores))
    return ler_entrada_matriz(matriz)


def ler_entrada(caminho: Path, aba: str = "Entrada") -> Tuple[List[Dict[str, Any]], List[str]]:
    ext = caminho.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return ler_xlsx_xlsm(caminho, aba)
    if ext == ".ods":
        return ler_ods(caminho, aba)
    raise ValueError("Formato não suportado. Use .xlsx, .xlsm ou .ods.")


def tela(nome: int, prox: int | None, numfields: int | None, numinputfields: int | None, entrada: str, entryscreen: bool = False) -> str:
    desc_extra = ""
    if numfields is not None:
        desc_extra += f'      <numfields number="{numfields}" optional="false" invertmatch="false" />\n'
    if numinputfields is not None:
        desc_extra += f'      <numinputfields number="{numinputfields}" optional="false" invertmatch="false" />\n'
    next_block = ""
    if prox is not None:
        next_block = f'''   <nextscreens timeout="0">
      <nextscreen name="Tela{prox}" />
   </nextscreens>
'''
    return f'''<screen name="Tela{nome}" entryscreen="{str(entryscreen).lower()}" exitscreen="false" transient="false">
   <description>
      <oia status="NOTINHIBITED" optional="false" invertmach="false" />
{desc_extra.rstrip()}
   </description>
   <actions>
      <input value="{entrada}" row="0" col="0" movecursor="true" xlatehostkeys="true" encrypted="false" />
   </actions>
{next_block}   </screen>
'''


def gerar_xml(registros: List[Dict[str, Any]], tecla_tipo: str = "44") -> str:
    partes: List[str] = []
    partes.append('<HAScript name="MACRO_MOVFIN" description="" timeout="600000" pausetime="300" promptball="true" blockinput="false" author="Gerador Python" creationdate="11/06/2026 00:00:00" supressclearevents="false" usevars="false" ignorepauseforenhancedtn="true" delayifnotenhancedtn="0" ignorepausetimeforenhancedtn="true">\n\n')
    total_telas = 1 + len(registros) * 7
    partes.append(tela(1, 2 if registros else None, None, None, "&gt;FPATMOVFIN[enter]", entryscreen=True))
    partes.append("\n")

    tela_atual = 2
    for i, item in enumerate(registros, start=1):
        siape = xml_attr(formatar_siape(item["SIAPE/7"]))
        rd = xml_attr(texto(item["R/D"]).lower())
        rubrica = xml_attr(formatar_rubrica(item["Rubrica"]))
        seq = xml_attr(formatar_seq(item["SEQ"]))
        mes = xml_attr(formatar_mes(item["MÊS/24"]))
        valor = xml_attr(formatar_valor(item["Valor/0000,00"]))
        justificativa = xml_attr(item["Justificativa"])
        doc = xml_attr(item["DOC. Legal"])
        nome = xml_attr(item.get("Nome", "")).replace("--", "- -")
        linha_excel = item.get("__linha_excel__", "")

        partes.append(f"<!-- Registro {i} | Linha Planilha {linha_excel} | SIAPE {siape} | {nome} -->\n")
        partes.append(tela(tela_atual, tela_atual + 1, 70, 5, f"{siape}[enter]"))
        partes.append("\n")
        partes.append(tela(tela_atual + 1, tela_atual + 2, 82, 6, f"{rd}{rubrica}{seq}i[tab]X[enter]"))
        partes.append("\n")
        partes.append(tela(tela_atual + 2, tela_atual + 3, 134, 17, f"{mes}{valor}[tab]{xml_attr(tecla_tipo)}[enter]"))
        partes.append("\n")
        partes.append(tela(tela_atual + 3, tela_atual + 4, 134, 7, f"{doc}[tab]{justificativa}[enter]"))
        partes.append("\n")
        partes.append(tela(tela_atual + 4, tela_atual + 5, 52, 1, "c[enter]"))
        partes.append("\n")
        partes.append(tela(tela_atual + 5, tela_atual + 6, 53, 0, "[enter]"))
        partes.append("\n")
        proxima = tela_atual + 7 if tela_atual + 6 < total_telas else None
        partes.append(tela(tela_atual + 6, proxima, 82, 6, "[pf12]"))
        partes.append("\n")
        tela_atual += 7

    partes.append("</HAScript>\n")
    return "".join(partes)


class App(tk.Frame):
    def __init__(self, master, planilha=None, aba="Entrada", bloco=None, rotulo=None, ao_inicio=None):
        super().__init__(master, bg='#f4f6f3')
        self.bloco = bloco
        self.rotulo = rotulo
        self.ao_inicio = ao_inicio
        self.pasta_saida = Path(planilha).parent if planilha else None
        self.controles = []
        self.pack(fill="both", expand=True)
        self._gerando = False

        self.planilha_var = tk.StringVar(value=str(planilha) if planilha else "")
        self.saida_var = tk.StringVar(value=str(self.pasta_saida / "macro.mac") if self.pasta_saida else "")
        self.aba_var = tk.StringVar(value=aba)
        self.tipo_var = tk.StringVar(value="44")
        self.status_var = tk.StringVar(value="Conferindo registros da planilha…" if planilha else "Escolha a planilha de entrada para começar.")

        self._montar_tela()
        if planilha:
            self.after_idle(self.conferir)

    def _texto(self, pai, texto='', tamanho=10, cor='#243529', negrito=False, **kwargs):
        if self.rotulo:
            return self.rotulo(pai, texto, tamanho, cor, negrito, **kwargs)
        return tk.Label(pai, text=texto, bg=pai.cget('bg'), fg=cor,
                        font=('Segoe UI', tamanho, 'bold' if negrito else 'normal'), **kwargs)

    def _cartao(self, expandir=False):
        if self.bloco:
            return self.bloco(self, expandir=expandir)
        frame = tk.Frame(self, bg='white', padx=20, pady=18)
        frame.pack(fill='both' if expandir else 'x', expand=expandir, pady=10)
        return frame

    def _montar_tela(self):
        self._texto(self, 'Gerador de Macro', 21, negrito=True).pack(pady=(0, 4))
        self._texto(self, 'Confira os lançamentos e gere o arquivo da macro.', cor='#66736a').pack(pady=(0, 8))
        cartao = self._cartao()
        self._texto(cartao, 'Dados da macro', 12, negrito=True).pack(anchor='w', pady=(0, 10))
        frame = tk.Frame(cartao, bg='white')
        frame.pack(fill='x')
        frame.columnconfigure(1, weight=1)
        for linha, titulo, var in [(0, 'Planilha de entrada', self.planilha_var), (1, 'Arquivo da macro', self.saida_var)]:
            self._texto(frame, titulo).grid(row=linha,column=0,sticky='w',padx=(0,12),pady=5)
            entrada=ttk.Entry(frame,textvariable=var,style='Caminho.TEntry',state='readonly')
            entrada.grid(row=linha,column=1,columnspan=3,sticky='ew',pady=5)
        self._texto(cartao, 'A macro será salva junto da planilha e do relatório de execução.', 9, '#66736a').pack(anchor='w',pady=(4,10))
        opcoes=tk.Frame(cartao,bg='white'); opcoes.pack(fill='x')
        self._texto(opcoes,'Aba da planilha').pack(side='left',padx=(0,8))
        aba=ttk.Entry(opcoes,textvariable=self.aba_var,width=17,style='Caminho.TEntry'); aba.pack(side='left')
        self._texto(opcoes,'Assunto de cálculo',negrito=True).pack(side='left',padx=(20,8))
        tipo=ttk.Entry(opcoes,textvariable=self.tipo_var,width=8,style='Caminho.TEntry'); tipo.pack(side='left')
        self.controles.extend([aba,tipo])
        acoes=tk.Frame(cartao,bg='white'); acoes.pack(fill='x',pady=(14,0))
        gerar=ttk.Button(acoes,text='Gerar macro  →',command=self.gerar,style='Acao.TButton')
        gerar.pack(side='right'); self.controles.append(gerar)
        aba.bind('<FocusOut>', lambda evento: self.conferir() if not self._gerando else None)
        aba.bind('<Return>', lambda evento: self.conferir() if not self._gerando else None)
        registros=self._cartao(expandir=True)
        self._texto(registros,'Registros da macro',12,negrito=True).pack(anchor='w',pady=(0,10))
        self._texto(registros,textvariable=self.status_var,cor='#66736a',wraplength=850,anchor='w').pack(fill='x',pady=(0,10))
        lista_frame=tk.Frame(registros,bg='#f1f4f1'); lista_frame.pack(fill='both',expand=True)
        self.lista=tk.Listbox(lista_frame,height=8,bg='#f1f4f1',fg='#36583d',font=('Consolas',10),
                             relief='flat',borderwidth=0,highlightthickness=0,selectbackground='#d2e4c8')
        scroll=ttk.Scrollbar(lista_frame,command=self.lista.yview)
        scroll.pack(side='right',fill='y'); self.lista.configure(yscrollcommand=scroll.set)
        self.lista.pack(fill='both',expand=True)
        self._texto(self,'UFGD  /  Divisão de Pagamento de Pessoal  •  progesp.dpp@ufgd.edu.br',9,'#66736a').pack(anchor='w',pady=(12,0))

    def _habilitar(self, ativo):
        for controle in self.controles:
            controle.configure(state='normal' if ativo else 'disabled')

    def _conclusao(self, total, ignoradas, saida):
        root=self.winfo_toplevel()
        dialogo=tk.Toplevel(root); dialogo.withdraw()
        dialogo.title('Macro gerada'); dialogo.transient(root)
        dialogo.configure(bg='white'); dialogo.resizable(False,False)
        painel=tk.Frame(dialogo,bg='white',padx=36,pady=28); painel.pack(fill='both',expand=True)
        self._texto(painel,'✓',36,'#567719',True).pack()
        self._texto(painel,'Macro gerada',22,negrito=True).pack(pady=(8,12))
        self._texto(painel,f'Registros usados: {total}   •   Linhas ignoradas: {ignoradas}',11,'#66736a').pack()
        self._texto(painel,str(saida),10,'#66736a',wraplength=500,justify='center').pack(pady=(14,20))
        botoes=tk.Frame(painel,bg='white'); botoes.pack()
        def voltar():
            dialogo.grab_release(); dialogo.destroy()
            if self.ao_inicio:
                self.ao_inicio()
            else:
                self.lista.delete(0,tk.END)
                self.status_var.set('Pronto para conferir registros.')
        def encerrar():
            dialogo.grab_release(); root.destroy()
        voltar_btn=ttk.Button(botoes,text='Retornar ao início',style='Acao.TButton',command=voltar)
        voltar_btn.pack(side='left',padx=(0,10))
        ttk.Button(botoes,text='Encerrar programa',style='Pasta.TButton',command=encerrar).pack(side='left')
        dialogo.update_idletasks()
        w,h=dialogo.winfo_reqwidth(),dialogo.winfo_reqheight()
        x=root.winfo_rootx()+(root.winfo_width()-w)//2; y=root.winfo_rooty()+(root.winfo_height()-h)//2
        dialogo.geometry(f'{w}x{h}+{max(0,x)}+{max(0,y)}')
        dialogo.protocol('WM_DELETE_WINDOW',voltar); dialogo.bind('<Return>',lambda e:voltar())
        dialogo.deiconify(); dialogo.grab_set(); voltar_btn.focus_set()

    def _validar_caminhos(self) -> Tuple[Path, Path, str, str] | None:
        planilha = Path(self.planilha_var.get().strip())
        pasta = self.pasta_saida or planilha.parent
        saida = pasta / "macro.mac"
        self.saida_var.set(str(saida))
        aba = self.aba_var.get().strip() or "Entrada"
        tipo = self.tipo_var.get().strip() or "44"
        if not planilha.is_file():
            messagebox.showerror("Erro", "Escolha uma planilha de entrada válida.")
            return None
        if planilha.suffix.lower() not in (".xlsx", ".xlsm", ".ods"):
            messagebox.showerror("Erro", "Formato não suportado. Use .xlsx, .xlsm ou .ods.")
            return None
        return planilha, saida, aba, tipo

    def conferir(self):
        self.lista.delete(0, tk.END)
        dados = self._validar_caminhos()
        if not dados:
            return
        planilha, _saida, aba, _tipo = dados
        try:
            registros, ignoradas = ler_entrada(planilha, aba=aba)
            self.lista.delete(0, tk.END)
            for i, item in enumerate(registros, start=1):
                self.lista.insert(tk.END, f"{i:03d} | Linha {item.get('__linha_excel__')} | SIAPE {formatar_siape(item['SIAPE/7'])} | {texto(item.get('Nome'))}")
            if ignoradas:
                self.lista.insert(tk.END, "--- Linhas ignoradas ---")
                for msg in ignoradas[:80]:
                    self.lista.insert(tk.END, msg)
            self.status_var.set(f"Conferência: {len(registros)} registros válidos. {len(ignoradas)} linhas ignoradas.")
        except Exception as e:
            self.status_var.set('Não foi possível conferir os registros. Verifique a planilha e a aba.')
            messagebox.showerror("Erro ao conferir", str(e))

    def gerar(self):
        if self._gerando:
            return
        dados = self._validar_caminhos()
        if not dados:
            return
        planilha, saida, aba, tipo = dados
        if saida.exists() and not messagebox.askyesno("Substituir macro", f"O arquivo já existe:\n{saida}\n\nDeseja substituí-lo?", parent=self):
            return
        self._gerando = True
        self._habilitar(False)
        self.status_var.set("Gerando macro…")

        def tarefa():
            try:
                registros, ignoradas = ler_entrada(planilha, aba=aba)
                xml = gerar_xml(registros, tecla_tipo=tipo)
                from xml.etree import ElementTree
                ElementTree.fromstring(xml)  # Impede gravar XML malformado.
                saida.parent.mkdir(parents=True, exist_ok=True)
                saida.write_text(xml, encoding="utf-8")
                self.after(0, lambda: self.sucesso(len(registros), len(ignoradas), saida))
            except Exception as e:
                self.after(0, lambda mensagem=str(e): self.erro(mensagem))

        threading.Thread(target=tarefa, daemon=True).start()

    def sucesso(self, total: int, ignoradas: int, saida: Path):
        self._gerando = False
        self._habilitar(True)
        self.saida_var.set(str(saida))
        self.status_var.set(f"Pronto: {total} registros processados. {ignoradas} linhas ignoradas. Macro: {saida}")
        self._conclusao(total, ignoradas, saida)

    def erro(self, mensagem: str):
        self._gerando = False
        self._habilitar(True)
        self.status_var.set("Erro ao gerar macro.")
        messagebox.showerror("Erro ao gerar macro", mensagem)


if __name__ == "__main__":
    root = tk.Tk()
    root.title("Gerador de Macro")
    root.geometry("980x600")
    App(root)
    root.mainloop()
